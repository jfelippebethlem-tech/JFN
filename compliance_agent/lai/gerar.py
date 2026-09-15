# -*- coding: utf-8 -*-
"""Um botão → o requerimento sai: fatos → texto → .docx/.md → registro com prazo."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from compliance_agent.lai import registro
from compliance_agent.lai.fatos import fatos_do_alvo
from compliance_agent.lai.requerimento import montar

_RAIZ = Path(__file__).resolve().parents[2]
_OUT = _RAIZ / "reports"


def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")[:60] or "alvo"


def _escrever(texto: str, nome_base: str) -> tuple[str, str]:
    """Mesmo padrão do mandato.py: .docx (python-docx) + .md legível/grep-ável para o Yoda."""
    _OUT.mkdir(parents=True, exist_ok=True)
    md = _OUT / f"{nome_base}.md"
    md.write_text(texto + "\n", encoding="utf-8")
    docx_path = _OUT / f"{nome_base}.docx"
    try:
        from docx import Document
        doc = Document()
        for i, ln in enumerate(texto.split("\n")):
            if i == 0:
                doc.add_heading(ln, level=0)
            elif ln.endswith(":") and ln.isupper():
                doc.add_heading(ln, level=2)
            else:
                doc.add_paragraph(ln)
        doc.add_paragraph("")
        doc.add_paragraph(f"Gerado pelo JFN em {datetime.now().strftime('%d/%m/%Y %H:%M')} — conferir o requerente e protocolar no e-SIC.")
        doc.save(str(docx_path))
    except ImportError:
        docx_path = None
    return (str(docx_path) if docx_path else None), str(md)


_PRIORIDADE = Path.home() / "shared-brain" / "sei_pcrj_prioridade.txt"


def _priorizar_no_sweep(numeros: list[str]) -> None:
    """Processo pedido na LAI sem árvore capturada vai para a FRENTE da fila da VM-2 (Syncthing →
    refresh_fila_prio.py). Nunca falha o requerimento por causa disso."""
    try:
        if not numeros:
            return
        atuais = set(_PRIORIDADE.read_text(encoding="utf-8").split()) if _PRIORIDADE.exists() else set()
        novos = [n for n in numeros if n not in atuais]
        if novos:
            _PRIORIDADE.parent.mkdir(parents=True, exist_ok=True)
            with open(_PRIORIDADE, "a", encoding="utf-8") as f:
                f.write("".join(n + "\n" for n in novos))
    except OSError:
        pass


def gerar(alvo: str, esfera: str | None = None, motivo: str | None = None, db_path=None, lai_db=None) -> dict:
    """Retorna {ok, id, alvo, esfera, destinatario, canal, prazo_dias, itens_pedido, path_docx, path_md,
    texto, resumo, avisos[], processos[], contratos[]}. Nunca levanta por base vazia: pede a íntegra."""
    alvo = (alvo or "").strip()
    if not alvo:
        return {"ok": False, "erro": "alvo vazio: informe nº de processo, contrato, CNPJ ou nome"}
    fatos = fatos_do_alvo(alvo, db_path)
    r = montar(fatos, esfera, motivo)
    from compliance_agent.reporting.neutralidade import garantir_neutro
    garantir_neutro(r["texto"], contexto="requerimento LAI")
    nome = f"lai_{_slug(alvo)}_{datetime.now().strftime('%Y%m%d_%H%M')}"
    path_docx, path_md = _escrever(r["texto"], nome)
    procs = [p["numero"] for p in fatos.get("processos") or []]
    contratos = [c["contrato"] for c in fatos.get("contratos") or []]
    id_ = registro.registrar(alvo, r["esfera"], r["destinatario"], procs, contratos, r["itens_pedido"],
                             path_docx, path_md, db_path=lai_db)
    docs_vistos = sum(len(p.get("documentos_vistos") or []) for p in fatos.get("processos") or [])
    _priorizar_no_sweep([p["numero"] for p in fatos.get("processos") or [] if not p.get("documentos_vistos")])
    resumo = (f"LAI #{id_} · {r['esfera']} · {len(procs)} processo(s), {len(contratos)} contrato(s), "
              f"{docs_vistos} documento(s) nomeados · {r['itens_pedido']} itens · prazo {r['prazo_dias']} dias após protocolo")
    return {"ok": True, "id": id_, "alvo": alvo, "esfera": r["esfera"], "destinatario": r["destinatario"], "canal": r["canal"],
            "base_legal": r["base_legal"], "prazo_dias": r["prazo_dias"], "itens_pedido": r["itens_pedido"],
            "path_docx": path_docx, "path_md": path_md, "texto": r["texto"], "resumo": resumo, "avisos": r["avisos"],
            "processos": procs, "contratos": contratos, "documentos_nomeados": docs_vistos}
