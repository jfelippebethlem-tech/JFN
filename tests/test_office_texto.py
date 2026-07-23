"""Extrair texto de anexos Office (Excel/Word) do SEI.

Achado 2026-07-23: anexos em formato Office (planilha de medição/faturamento EXCEL,
minuta Word) faziam `_conteudo_via_arvore` cair em `tipo=None → continue` → doc
perdido. São AUDIT-CRÍTICOS (a planilha tem os números da medição). Tooling presente:
openpyxl (.xlsx), xlrd (.xls), python-docx (.docx). Só o .doc binário antigo fica de
fora (precisa antiword/libreoffice, não instalados).
"""
import io

from compliance_agent.sei.office_texto import texto_de_office


def _xlsx_bytes(linhas):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for ln in linhas:
        ws.append(ln)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(paragrafos):
    import docx
    d = docx.Document()
    for p in paragrafos:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_extrai_planilha_xlsx():
    body = _xlsx_bytes([["Item", "Valor"],
                        ["Aluguel container", "12.500,00"],
                        ["3a Medição", "TOTAL 250.000,00"]])
    txt = texto_de_office(body)
    assert "Aluguel container" in txt
    assert "250.000,00" in txt
    assert "Medição" in txt


def test_extrai_word_docx():
    body = _docx_bytes(["MINUTA DE RESOLUÇÃO SES",
                        "Art. 1º Fica aprovada a contratação.",
                        "Valor global: R$ 1.200.000,00"])
    txt = texto_de_office(body)
    assert "MINUTA DE RESOLUÇÃO" in txt
    assert "1.200.000,00" in txt


def test_doc_binario_antigo_devolve_vazio_sem_quebrar():
    """.doc OLE2 antigo (sem antiword): devolve '' honesto, não explode."""
    ole2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 200
    assert texto_de_office(ole2) == ""


def test_bytes_lixo_devolve_vazio():
    assert texto_de_office(b"nao e office nenhum") == ""
    assert texto_de_office(b"") == ""


def test_pdf_nao_e_tratado_aqui():
    """PDF não é Office — devolve '' (o caminho de PDF/OCR cuida dele)."""
    assert texto_de_office(b"%PDF-1.7 ...") == ""
